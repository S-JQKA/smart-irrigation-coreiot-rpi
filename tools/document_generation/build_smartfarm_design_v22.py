from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


WORKSPACE = Path(__file__).resolve().parents[2]
TEMPLATE = Path(r"C:\Users\voles\Downloads\Thiet_ke_he_thong_v2.1_SmartFarm_an_toan.docx")
OUTPUT = (
    WORKSPACE
    / "docs"
    / "final"
    / "system_design"
    / "RB Thiet_ke_he_thong_v2.2_SmartFarm_an_toan_rebuild.docx"
)

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(11, 37, 69)
MUTED = RGBColor(95, 95, 95)
BLACK = RGBColor(0, 0, 0)
HEADER_FILL = "E8EEF5"
LIGHT_FILL = "F4F6F9"
WARN_FILL = "FFF4D6"
BORDER = "B7C4D6"
WHITE = "FFFFFF"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def clear_body(doc: Document) -> None:
    body = doc._body._element
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    if run._element.rPr is None:
        run._element.get_or_add_rPr()
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_spacing(p, before=0, after=6, line=1.10):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def configure_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        st = styles[name]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number", "List Paragraph"):
        if name in styles:
            st = styles[name]
            st.font.name = "Calibri"
            st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
            st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
            st.font.size = Pt(11)
            st.paragraph_format.space_after = Pt(4)
            st.paragraph_format.line_spacing = 1.167


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def configure_header_footer(doc: Document) -> None:
    section = doc.sections[0]
    hp = section.header.paragraphs[0]
    hp.clear()
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = hp.add_run("Smart Farm · ESP-NOW · Central/Manifold · Raspberry Pi Gateway · CoreIoT")
    set_run_font(r, size=9, color=MUTED)

    fp = section.footer.paragraphs[0]
    fp.clear()
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = fp.add_run("HK53-DATN-008 — Thiết kế hệ thống v2.2\tTrang ")
    set_run_font(r, size=9, color=MUTED)
    add_page_field(fp)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, bottom=80, start=120, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        elem = tc_mar.find(qn(f"w:{m}"))
        if elem is None:
            elem = OxmlElement(f"w:{m}")
            tc_mar.append(elem)
        elem.set(qn("w:w"), str(v))
        elem.set(qn("w:type"), "dxa")


def set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        tbl.remove(grid)
    grid = OxmlElement("w:tblGrid")
    for width in widths:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(width))
        grid.append(gc)
    tbl.insert(0, grid)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths[min(idx, len(widths) - 1)]
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table, color=BORDER):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = borders.find(qn(f"w:{edge}"))
        if elem is None:
            elem = OxmlElement(f"w:{edge}")
            borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), "4")
        elem.set(qn("w:space"), "0")
        elem.set(qn("w:color"), color)


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def format_cell(cell, bold=False, size=9.2, color=BLACK, align=None):
    for p in cell.paragraphs:
        if align is not None:
            p.alignment = align
        set_paragraph_spacing(p, before=0, after=0, line=1.05)
        for run in p.runs:
            set_run_font(run, size=size, color=color, bold=bold)


def add_para(doc, text="", bold=False, italic=False, color=BLACK, size=11, align=None, before=0, after=6):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=before, after=after)
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    set_run_font(r, size=size, color=color, bold=bold, italic=italic)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)
        for run in p.runs:
            set_run_font(run, size=11)


def add_numbers(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)
        for run in p.runs:
            set_run_font(run, size=11)


def add_table(doc, headers, rows, widths=None, font_size=9.0, header_fill=HEADER_FILL):
    if widths is None:
        widths = [TABLE_WIDTH_DXA // len(headers)] * len(headers)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    set_table_borders(table)
    mark_header_row(table.rows[0])

    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        shade_cell(cell, header_fill)
        format_cell(cell, bold=True, size=font_size, color=INK, align=WD_ALIGN_PARAGRAPH.CENTER)

    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            shade_cell(cells[i], WHITE)
            align = WD_ALIGN_PARAGRAPH.CENTER if len(str(value)) < 18 and i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            format_cell(cells[i], size=font_size, align=align)

    add_para(doc, "", after=4)
    return table


def add_callout(doc, title, body, fill=LIGHT_FILL):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [TABLE_WIDTH_DXA])
    set_table_borders(table, "C9D3E2")
    mark_header_row(table.rows[0])
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    p = cell.paragraphs[0]
    p.text = ""
    set_paragraph_spacing(p, after=3)
    r = p.add_run(title)
    set_run_font(r, size=10.5, bold=True, color=DARK_BLUE)
    p2 = cell.add_paragraph()
    set_paragraph_spacing(p2, after=0, line=1.10)
    r2 = p2.add_run(body)
    set_run_font(r2, size=10.2, color=BLACK)
    add_para(doc, "", after=4)
    return table


def add_code_block(doc, text):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [TABLE_WIDTH_DXA])
    set_table_borders(table, "D8DDE6")
    mark_header_row(table.rows[0])
    cell = table.cell(0, 0)
    shade_cell(cell, "F7F9FB")
    cell.text = text
    for p in cell.paragraphs:
        set_paragraph_spacing(p, after=0, line=1.0)
        for run in p.runs:
            set_run_font(run, name="Consolas", size=8.7, color=INK)
    add_para(doc, "", after=4)


def add_metadata(doc):
    rows = [
        ("Mã đề tài", "HK53-DATN-008"),
        ("GVHD", "TS. Lê Trọng Nhân"),
        ("Sinh viên", "Võ Lê Sinh — MSSV 2212927"),
        ("Bộ môn / chương trình", "Kỹ thuật Máy tính · Tiếng Việt · Hướng ứng dụng"),
        ("Phiên bản tài liệu", "v2.2 rebuild — cập nhật từ SRS v2.2 và draft thiết kế v2.2"),
        ("Trạng thái", "Bản thiết kế hệ thống dùng cho chốt phương án, triển khai và kiểm thử"),
    ]
    add_table(doc, ["Hạng mục", "Thông tin"], rows, widths=[2100, 7260], font_size=9.6)
    add_callout(
        doc,
        "Phạm vi sử dụng",
        "Tài liệu này thay thế bản thiết kế v2.1 ở phần kiến trúc nước/flow, mode điều khiển, scheduler và safety mở rộng. "
        "Tinh thần trình bày giữ theo bản v2.1: mỗi quyết định kiến trúc quan trọng đều có lý do, đánh đổi và tiêu chí kiểm chứng.",
    )


def build_doc():
    # The v2.1 file is a visual reference, but its style tree is intentionally
    # minimal and lacks standard Word styles such as Normal/List Bullet/Table Grid.
    # Build from a clean DOCX and recreate the v2.1 presentation patterns.
    doc = Document()
    configure_styles(doc)
    configure_header_footer(doc)

    add_para(doc, "TÀI LIỆU THIẾT KẾ HỆ THỐNG", bold=True, size=22, color=INK, after=4)
    add_para(doc, "Phiên bản 2.2 — Smart Farm đa zone an toàn trên nền Smart Irrigation Template", size=13.5, color=DARK_BLUE, after=3)
    add_para(doc, "Main pump + manifold · valve/flow theo zone · ESP-NOW · Raspberry Pi Gateway · CoreIoT", size=11.5, color=MUTED, after=12)
    add_metadata(doc)
    add_para(doc, "Mục lục", bold=True, size=14, color=INK, before=4, after=5)
    add_bullets(
        doc,
        [
            "0. Thay đổi so với phiên bản 2.1",
            "1. Phạm vi và nguyên tắc thiết kế",
            "2. Kiến trúc tổng thể",
            "3. Mạng ESP-NOW, sensor node và Central/Manifold Controller",
            "4. Mô hình thực thể trên CoreIoT",
            "5. Hợp đồng dữ liệu và ánh xạ telemetry",
            "6. Topology nước, flow sensor theo zone và thống kê nước",
            "7. Logic điều khiển, Irrigation Scheduler và mode vận hành",
            "8. Cảnh báo, WaterloggingRisk và drainage hook",
            "9. Thiết kế phần mềm Raspberry Pi gateway",
            "10. Module phát hiện bất thường",
            "11. Thiết kế an toàn",
            "12. Dashboard giám sát và vận hành",
            "13. Thiết kế phần cứng",
            "14. Kiểm thử, nghiệm thu và lộ trình",
        ],
    )

    add_heading(doc, "0. Thay đổi so với phiên bản 2.1", 1)
    add_para(
        doc,
        "Bản 2.1 đã đúng ở định hướng Smart Farm đa cảm biến, Raspberry Pi gateway, ESP-NOW và safety nhiều lớp. "
        "Tuy nhiên phần nước/flow vẫn còn gần với mô hình demo: bơm/van điều khiển tập trung, water meter mở và khả năng tưới song song chưa rõ. "
        "Bản 2.2 chuyển sang topology có thể mở rộng theo zone, đồng thời giữ các nguyên tắc an toàn đã chốt ở v2.1.",
    )
    add_table(
        doc,
        ["Nhóm thay đổi", "Thiết kế v2.1", "Thiết kế v2.2 mới"],
        [
            ("Topology nước", "Bơm/actuator được mô tả chung; water meter có thể mô phỏng hoặc thay bằng YF-S201.", "Một main pump cấp manifold; mỗi zone có valve và flow sensor nhánh riêng ở reference design."),
            ("Điều khiển actuator", "Pi có thể điều khiển relay trực tiếp trong mô hình lab.", "Reference topology dùng ESP32 Central/Manifold Controller đặt gần trạm nước; Pi chỉ điều phối lệnh."),
            ("Flow / water statistic", "Một `pulseCounter` lõi, mô phỏng trước rồi thay driver.", "Mỗi zone có `pulseCounter_Zi`; khi publish vào logical `SI Water Meter` của zone vẫn giữ key lõi `pulseCounter`."),
            ("Tưới song song", "Chưa mô tả scheduler rõ.", "Có `Irrigation Scheduler`, giới hạn `maxConcurrentZones`, đóng từng zone độc lập khi đạt target/hạn mức."),
            ("Mực nước", "`waterLevel` vừa là telemetry vừa dùng cho safety.", "Tách `tankLowSwitch`/phao cạn hard interlock (Must) khỏi `waterLevel`/`tankLevelPct` hiển thị (Should)."),
            ("Tiêu / chống úng", "High moisture và hạn mức nước đã có nhưng chưa thành lớp tiêu tối thiểu.", "Thêm `WaterloggingRisk`, khóa tưới zone quá ẩm; Drain Pump/Valve là hook mở rộng."),
            ("pH", "Được nêu như tùy chọn trong Env Sensor Cluster.", "Đưa khỏi lõi hiện tại; chỉ giữ future work cho nutrient/fertigation."),
            ("Data quality", "Drop/clamp/gắn cờ ở node/gateway.", "Không publish giá trị rác vào key chính; publish quality/status/last-invalid để giải thích alarm."),
            ("Mode", "NORMAL/DEGRADED/SAFE-IDLE đã có.", "Tách rõ Control mode `MANUAL/AUTO/DISABLED` theo field khỏi System mode `NORMAL/DEGRADED/SAFE-IDLE`."),
            ("Kiểm thử", "37 test case.", "45 test case, bổ sung topology branch flow, scheduler, Central Controller, WaterloggingRisk, manual TTL/latch."),
        ],
        widths=[1800, 3400, 4160],
        font_size=8.7,
    )
    add_callout(
        doc,
        "Quyết định #1 — v2.2 là bản đổi topology, không chỉ chỉnh từ ngữ",
        "Giá trị mới của bản 2.2 nằm ở việc đưa phần nước/flow từ mô hình demo sang mô hình nhiều zone có thể mở rộng: "
        "main pump + manifold + valve/flow từng nhánh + controller tại trạm nước. Các lớp cloud, data contract và safety của v2.1 được giữ và làm rõ.",
    )

    add_heading(doc, "1. Phạm vi và nguyên tắc thiết kế", 1)
    add_heading(doc, "1.1. Bài toán", 2)
    add_para(
        doc,
        "Hệ thống cần vận hành thử nghiệm một Smart Farm đa cảm biến - đa cơ cấu chấp hành trên CoreIoT, lấy Smart Irrigation Template làm lõi nghiệp vụ. "
        "Mỗi zone có cụm cảm biến môi trường riêng, dữ liệu được truyền không dây bằng ESP-NOW về gateway, sau đó ánh xạ thành các logical device trên CoreIoT. "
        "Phần trạm nước dùng một main pump, manifold, valve và flow sensor theo từng zone để vừa hỗ trợ tưới song song vừa thống kê nước đúng cho từng `SI Field`.",
    )
    add_heading(doc, "1.2. Bảy nguyên tắc xuyên suốt", 2)
    add_bullets(
        doc,
        [
            "Template là lõi tùy biến: giữ đúng key/RPC lõi, nhưng được mở rộng device profile, alarm, rule chain, widget.",
            "Pi là edge gateway điều phối, mapping, cache, replay và safety cấp hệ thống; Pi không mặc định kéo dây relay xa.",
            "ESP32 tại hiện trường làm đúng việc hiện trường: sensor node đo môi trường, Central/Manifold Controller đọc flow/low-water và điều khiển actuator cục bộ.",
            "Safety đứng ngang hàng chức năng: hard interlock luôn thắng mọi lệnh cloud/manual/auto.",
            "Mất cloud không đồng nghĩa tắt mù: chuyển DEGRADED và chạy local rule tối giản trong giới hạn cứng.",
            "Dữ liệu xấu không được làm bẩn time-series chính; nhưng phải để lại quality/status đủ giải thích cảnh báo.",
            "pH, fertigation, drainage chủ động và OTA đầy đủ là hướng mở rộng; không được chặn tiến độ lõi tưới tiêu.",
        ],
    )
    add_heading(doc, "1.3. Trong / ngoài phạm vi", 2)
    add_table(
        doc,
        ["Trong phạm vi lõi v2.2", "Ngoài phạm vi hoặc thiết kế mở"],
        [
            ("Đa cảm biến: moisture, airTemp, airHumidity, lightLux, tankLowSwitch, waterLevel/tankLevelPct tùy chọn.", "pH liên tục dài ngày; nutrient/fertigation; cảm biến dinh dưỡng chuyên sâu."),
            ("ESP-NOW nhiều sensor node và Central/Manifold Controller qua bridge/receiver.", "LoRaWAN, cellular fallback, multi-tenant production."),
            ("Main pump + manifold + valve/flow sensor theo zone; SIM mode được phép nếu giữ interface.", "Drain pump/valve chạy thật và thuật toán tiêu nước đầy đủ."),
            ("CoreIoT Smart Irrigation Template, Direct Device API, đường chuyển Gateway API.", "Tùy biến template ngoài phạm vi dashboard/rule/alarm cần cho đồ án."),
            ("Safety: phao cạn hard interlock, watchdog, dry-run/low-flow, hạn mức nước, waterlogging lock, fail-safe relay.", "Chứng nhận an toàn điện công nghiệp hoặc vận hành ngoài trời dài hạn."),
            ("Anomaly Detector mức cơ bản: z-score/EWMA/rule/Isolation Forest nhẹ.", "ML mức C/D như dự báo dài hạn, RL tối ưu tưới hoặc mô hình cần dữ liệu nhãn lớn."),
            ("OTA partition/versioning ở mức thiết kế mở; cập nhật qua USB trong phạm vi demo.", "OTA production end-to-end qua cloud."),
        ],
        widths=[4680, 4680],
        font_size=8.8,
    )

    add_heading(doc, "2. Kiến trúc tổng thể", 1)
    add_callout(
        doc,
        "Quyết định #2 — Phân vai bốn tầng",
        "Kiến trúc v2.2 tách rõ: sensor node theo zone sinh dữ liệu; Central/Manifold Controller thực thi nước/flow gần trạm; Raspberry Pi chuẩn hóa và điều phối; CoreIoT giữ dashboard, rule chain, alarm và lưu trữ.",
    )
    add_table(
        doc,
        ["Tầng", "Thành phần", "Trách nhiệm"],
        [
            ("Hiện trường - zone", "ESP32 sensor node + soil capacitive + SHT31/DHT22 + BH1750 + pin/RSSI.", "Đọc môi trường từng zone, kiểm dữ liệu sơ cấp, gửi ESP-NOW uplink."),
            ("Hiện trường - trạm nước", "ESP32 Central/Manifold Controller, main pump relay/driver, valve từng zone, flow sensor nhánh, low-water switch.", "Điều khiển actuator tại chỗ, đếm flow từng nhánh, enforce hard interlock trước khi bật bơm/valve."),
            ("Bridge/receiver", "ESP32 ESP-NOW bridge/receiver nối Raspberry Pi qua UART.", "Nhận uplink từ sensor/Central; gửi downlink command tới Central; giữ kênh radio cố định."),
            ("Edge gateway", "Raspberry Pi 4, 11 module phần mềm, cache cấu hình, replay buffer.", "Parse, quality gate, mapping logical device, MQTT QoS 1, scheduler, local safety, anomaly, log/replay."),
            ("Cloud/CoreIoT", "Smart Irrigation Template + profile mở rộng + rule chain + dashboard + alarm.", "Lưu time-series, hiển thị, rule cloud, RPC, alarm routing, cấu hình attribute."),
        ],
        widths=[1450, 3600, 4310],
        font_size=8.8,
    )
    add_code_block(
        doc,
        "Zone sensor nodes --ESP-NOW--> ESP32 bridge --UART--> Raspberry Pi --MQTT--> CoreIoT\n"
        "CoreIoT RPC/attribute --> Raspberry Pi scheduler/safety --> ESP-NOW command --> Central/Manifold Controller\n"
        "Central Controller --> valve/main pump/flow/low-water telemetry --> bridge --> Pi --> CoreIoT logical devices"
    )
    add_table(
        doc,
        ["Luồng", "Mô tả", "Điểm kiểm soát"],
        [
            ("Telemetry môi trường", "Node gửi moisture/T/RH/lux/battery/RSSI/sequence.", "CRC, MAC whitelist, range/rate gate, timestamp tại Pi."),
            ("Telemetry nước", "Central gửi tankLowSwitch, valve state, pump state, pulseCounter_Zi, flow_Zi.", "Monotonic counter, branch mapping, low-water hard interlock."),
            ("RPC/manual", "Dashboard gửi TURN_ON/OFF hoặc mode/attribute; Pi kiểm whitelist và metadata.", "commandId, timestamp, TTL, safety priority, ACK từ Central."),
            ("Auto/scheduler", "Pi hoặc cloud rule chọn zone cần tưới, không vượt maxConcurrentZones.", "Hysteresis, debounce, maxDuration, maxWaterPerCycle/day."),
            ("Mất cloud", "Pi chuyển DEGRADED sau ngưỡng mất MQTT.", "Local Control Engine dùng cache; replay khi NORMAL trở lại."),
        ],
        widths=[1600, 4200, 3560],
        font_size=8.8,
    )

    add_heading(doc, "3. Mạng ESP-NOW, sensor node và Central/Manifold Controller", 1)
    add_heading(doc, "3.1. Sensor node theo zone", 2)
    add_para(
        doc,
        "Mỗi zone có một ESP32 sensor node độc lập. Node không cần biết logic tưới đầy đủ; node chỉ đo, kiểm lỗi đọc cảm biến cơ bản, đóng gói frame và gửi về bridge. "
        "Việc quyết định tưới nằm ở Pi/CoreIoT để tránh phân tán nghiệp vụ và khó kiểm chứng.",
    )
    add_table(
        doc,
        ["Trường frame", "Ý nghĩa", "Ghi chú"],
        [
            ("source_mac", "MAC node gửi.", "Dùng ánh xạ sang logical device trong `nodeMap.yaml`."),
            ("node_type", "`sensor_node` hoặc `central_controller`.", "Giúp parser chọn schema."),
            ("seq", "Số thứ tự gói.", "Phát hiện mất gói, trùng gói, replay."),
            ("ts_node", "Timestamp tương đối nếu có.", "Timestamp quyết định vẫn do Pi gắn khi nhận."),
            ("payload", "Danh sách key/value.", "moisture, airTemp, airHumidity, lightLux, battery, rssi..."),
            ("status", "Cờ lỗi cảm biến/CRC/I2C/low battery.", "Không publish rác vào key chính."),
            ("crc", "Checksum frame.", "Sai CRC bị drop tại receiver/Pi."),
        ],
        widths=[1700, 3300, 4360],
        font_size=8.7,
    )
    add_heading(doc, "3.2. ESP32 bridge/receiver/transceiver", 2)
    add_para(
        doc,
        "Bridge chạy ESP-NOW ở kênh cố định, nối Pi qua UART 115200 8N1. Ở v2.2 bridge không chỉ nhận uplink mà còn chuyển downlink command từ Pi tới Central/Manifold Controller. "
        "Vì vậy frame cần có direction, message type, sequence, ACK/NAK và reason code.",
    )
    add_heading(doc, "3.3. Central/Manifold Controller", 2)
    add_table(
        doc,
        ["Nhiệm vụ", "Mô tả thiết kế"],
        [
            ("Đặt gần trạm nước", "Controller nằm gần bơm/manifold để dây relay, flow pulse và low-water switch ngắn, giảm nhiễu và rủi ro cơ khí."),
            ("Quản lý 4-6 zone/cụm", "Một controller đọc nhiều flow nhánh và điều khiển nhiều valve; hệ 10-20 zone mở rộng bằng nhiều controller."),
            ("Hard interlock cục bộ", "Nếu phao cạn kích hoạt, controller không bật main pump dù Pi/cloud yêu cầu."),
            ("ACK command", "Mọi lệnh actuator có commandId, TTL, trạng thái thực thi và reason khi bị block."),
            ("Telemetry trạng thái", "Gửi pumpState, valveState_Zi, pulseCounter_Zi, flow_Zi, tankLowSwitch, fault flags."),
        ],
        widths=[2200, 7160],
        font_size=8.8,
    )
    add_heading(doc, "3.4. Đánh đổi truyền thông", 2)
    add_table(
        doc,
        ["Tiêu chí", "ESP-NOW (chọn)", "Wi-Fi/MQTT trực tiếp", "LoRa"],
        [
            ("Hạ tầng", "Không cần AP, peer-to-peer.", "Cần AP/router ổn định.", "Cần gateway/server LoRaWAN nếu đi chuẩn."),
            ("Độ trễ", "Mức ms, hợp điều khiển gần thời gian thực.", "50-500 ms, phụ thuộc AP.", "1-5 s, bị duty cycle."),
            ("Tầm xa thực tế", "Tốt cho farm vài chục đến ~100 m; có thể tối ưu anten/vị trí.", "Phụ thuộc AP và vật cản.", "Rất xa nhưng datarate thấp."),
            ("Năng lượng", "Thấp hơn Wi-Fi connected.", "Cao nếu giữ MQTT/keepalive.", "Thấp nhất."),
            ("OTA", "Khó hơn, để thiết kế mở.", "Dễ nhất.", "Không phù hợp firmware lớn."),
            ("Kết luận", "Chọn cho node và Central vì ít hạ tầng, trễ thấp.", "Không chọn làm mặc định.", "Không chọn cho điều khiển thời gian thực của demo."),
        ],
        widths=[1600, 2900, 2500, 2360],
        font_size=8.4,
    )
    add_callout(
        doc,
        "Quyết định #3 — Central Controller cũng đi qua ESP-NOW",
        "Giữ chung một lớp truyền cho sensor node và trạm nước giúp đơn giản hóa gateway: Pi chỉ cần một bridge/receiver, một parser và một cơ chế ACK/retry. "
        "Khi cần triển khai xa hơn, có thể thêm nhiều bridge hoặc chuyển riêng trạm nước sang RS485/MQTT mà không phá contract CoreIoT.",
    )

    add_heading(doc, "4. Mô hình thực thể trên CoreIoT", 1)
    add_heading(doc, "4.1. Lõi Smart Irrigation Template giữ nguyên", 2)
    add_table(
        doc,
        ["Thực thể lõi", "Loại", "Dữ liệu / vai trò"],
        [
            ("SI Field", "Asset", "Một zone tưới; chứa sensor, water meter và valve/controller; có averageMoisture/waterConsumption."),
            ("SI Soil Moisture Sensor", "Device", "Telemetry `moisture`, `battery`; tham gia averageMoisture và alarm low/high moisture."),
            ("SI Water Meter", "Device", "Telemetry lõi `pulseCounter`, `battery`; mỗi zone có một logical meter."),
            ("SI Smart Valve", "Device", "RPC `TURN_ON`/`TURN_OFF`; trong v2.2 đại diện valve zone hoặc pattern cho Valve Controller."),
            ("Dashboard/rule/alarm lõi", "Template", "Giữ contract lõi để tái dùng dashboard, calculated field, alarm và lịch tưới."),
        ],
        widths=[2200, 1300, 5860],
        font_size=8.8,
    )
    add_heading(doc, "4.2. Phần mở rộng có kỷ luật", 2)
    add_table(
        doc,
        ["Thực thể mở rộng", "Vai trò", "Ghi chú tích hợp"],
        [
            ("Env Sensor Cluster", "airTemp, airHumidity, lightLux, waterLevel/tankLevelPct, battery, rssi.", "pH không nằm trong lõi hiện tại; chỉ future work."),
            ("Pump Controller", "Main pump state/RPC.", "Main pump chỉ ON khi có ít nhất một valve active và qua interlock."),
            ("Valve Controller", "Valve từng zone, có thể dùng pattern `SI Smart Valve`.", "Mỗi zone có valve state riêng."),
            ("Central/Manifold Controller", "Logical device hoặc nhóm telemetry cho trạm nước.", "Gửi ACK, fault, tankLowSwitch, branch flow."),
            ("Drain Controller", "Hook cho drainPump/drainValve.", "Should/Future; chưa bắt buộc thuật toán tiêu chủ động."),
            ("Rule chain mở rộng", "Multivar control, scheduler hook, alarm routing.", "Không sửa contract lõi, chỉ thêm nhánh xử lý."),
        ],
        widths=[2300, 3300, 3760],
        font_size=8.8,
    )
    add_callout(
        doc,
        "Quyết định #4 — Logical device theo zone, physical controller có thể gom nhiều zone",
        "Một ESP32 Central/Manifold Controller vật lý có thể quản lý 4-6 zone, nhưng trên CoreIoT vẫn nên biểu diễn theo logical device của từng zone để dashboard, alarm và water statistic không bị trộn.",
    )

    add_heading(doc, "5. Hợp đồng dữ liệu và ánh xạ telemetry", 1)
    add_heading(doc, "5.1. Hai lớp key", 2)
    add_table(
        doc,
        ["Lớp", "Key / method", "Nguyên tắc"],
        [
            ("Lõi template", "`moisture`, `pulseCounter`, `battery`, `TURN_ON`, `TURN_OFF`.", "Không đổi tên khi publish vào logical device lõi."),
            ("Mở rộng môi trường", "`airTemp`, `airHumidity`, `lightLux`, `rssi`, `fw_version`.", "Publish vào Env Sensor Cluster hoặc asset field theo mapping."),
            ("Mở rộng nước", "Nội bộ: `pulseCounter_Zi`, `flow_Zi`, `valveState_Zi`, `tankLowSwitch`.", "Khi map sang `SI Water Meter Zi`, key public vẫn là `pulseCounter`."),
            ("Quality/status", "`quality_<key>`, `lastInvalid_<key>`, `dropReason`, `seqGap`.", "Giải thích alarm mà không làm bẩn chart chính."),
            ("Command metadata", "`commandId`, `source`, `issuedAt`, `ttlMs`, `mode`, `reason`.", "Chống lệnh cũ, retry trùng và thao tác không truy vết."),
        ],
        widths=[1750, 3800, 3810],
        font_size=8.6,
    )
    add_heading(doc, "5.2. Mapping theo zone", 2)
    add_table(
        doc,
        ["Nguồn vật lý", "Gateway internal key", "Logical device CoreIoT", "Key publish"],
        [
            ("Sensor node Z1", "moisture_Z1, airTemp_Z1, lightLux_Z1", "SI Soil Moisture Sensor Z1 / Env Sensor Cluster Z1", "moisture, airTemp, lightLux"),
            ("Flow branch Z1", "pulseCounter_Z1, flow_Z1", "SI Water Meter Z1", "pulseCounter"),
            ("Valve branch Z1", "valveState_Z1", "SI Smart Valve / Valve Controller Z1", "state / telemetry mở rộng"),
            ("Tank low switch", "tankLowSwitch", "Pump Controller / Central Controller", "tankLowSwitch"),
            ("Central Controller", "pumpState, controllerAck, faultFlags", "Pump Controller / Central Controller", "pumpState, ackStatus, faultFlags"),
        ],
        widths=[2300, 2450, 2700, 1910],
        font_size=8.3,
    )
    add_heading(doc, "5.3. Payload mẫu", 2)
    add_code_block(
        doc,
        "Sensor node Z1 -> Pi internal:\n"
        "{mac:'A0:B1:C2:01', node_type:'sensor_node', seq:1281,\n"
        " payload:{moisture:24.8, airTemp:34.1, airHumidity:51.2, lightLux:42600, battery:82, rssi:-68}}\n\n"
        "Central Controller -> Pi internal:\n"
        "{mac:'A0:B1:C2:CC', node_type:'central_controller', seq:942,\n"
        " payload:{pumpState:1, tankLowSwitch:false, valveState_Z1:1, pulseCounter_Z1:23840, flow_Z1:2.1}}\n\n"
        "Pi -> SI Water Meter Z1 telemetry:\n"
        "{pulseCounter:23840, battery:100}"
    )
    add_heading(doc, "5.4. Direct Device API và đường chuyển Gateway API", 2)
    add_para(
        doc,
        "Bản triển khai đầu dùng Direct Device API: mỗi logical device có access token riêng, Pi publish MQTT `v1/devices/me/telemetry` theo token tương ứng. "
        "Thiết kế vẫn giữ adapter để chuyển sang Gateway API `v1/gateway/telemetry` khi số device tăng hoặc muốn giảm số connection.",
    )
    add_callout(
        doc,
        "Quyết định #5 — Không publish `pulseCounter_Zi` trực tiếp vào key lõi",
        "Trong nội bộ gateway, suffix `_Zi` giúp phân biệt branch. Khi dữ liệu đi vào từng logical `SI Water Meter`, key phải quay về `pulseCounter` để giữ đúng Smart Irrigation Template và calculated field `waterConsumption`.",
    )

    add_heading(doc, "6. Topology nước, flow sensor theo zone và thống kê nước", 1)
    add_heading(doc, "6.1. Reference topology", 2)
    add_table(
        doc,
        ["Thành phần", "Vai trò", "Ghi chú thiết kế"],
        [
            ("Main pump", "Cấp áp/lưu lượng cho manifold.", "ON khi có ít nhất một zone active; OFF khi không còn zone active hoặc safety block."),
            ("Manifold", "Chia nước sang các nhánh zone.", "Nên đặt gần controller và flow sensor để dây ngắn."),
            ("Valve Zi", "Mở/đóng tưới riêng từng zone.", "Scheduler có thể mở nhiều valve trong giới hạn `maxConcurrentZones`."),
            ("Flow sensor Zi", "Đếm nước riêng từng nhánh.", "Dùng YF-S201 hoặc SIM mode cùng interface."),
            ("Low-water switch", "Hard interlock bể cạn.", "Tín hiệu safety Must, không thay bằng đo liên tục."),
            ("Tank level sensor", "Hiển thị mức nước phần trăm/ước lượng.", "Should, không được vượt quyền phao cạn."),
        ],
        widths=[2000, 3000, 4360],
        font_size=8.8,
    )
    add_heading(doc, "6.2. Vì sao không dùng một flow tổng duy nhất", 2)
    add_table(
        doc,
        ["Một flow tổng", "Flow riêng từng nhánh"],
        [
            ("Rẻ hơn và ít wiring hơn.", "Tốn linh kiện hơn nhưng thống kê đúng theo từng zone."),
            ("Chỉ chính xác nếu tưới một zone tại một thời điểm.", "Cho phép tưới song song mà vẫn tính `waterConsumption` riêng."),
            ("Nếu tưới song song phải chia ảo theo thời gian/áp lực, dễ sai.", "Dry-run/ZoneFlowLow/ValveLeak bắt được theo từng branch."),
            ("Phù hợp demo đơn giản.", "Phù hợp thiết kế v2.2 có scheduler và mở rộng 10-20 zone."),
        ],
        widths=[4680, 4680],
        font_size=8.8,
    )
    add_heading(doc, "6.3. Công thức thống kê nước", 2)
    add_para(
        doc,
        "`waterConsumption_Zi` được tính từ delta pulse của branch Zi: `liters = deltaPulse / pulsesPerLiter`. "
        "Với YF-S201 thường dùng xấp xỉ 450 pulse/L sau hiệu chuẩn; biến `pulsesPerLiter` phải để trong config theo từng flow sensor. "
        "Pi lưu daily volume và cycle volume để enforce `maxWaterPerCycle` và `maxWaterPerDay`.",
    )
    add_heading(doc, "6.4. Mở rộng nhiều zone", 2)
    add_table(
        doc,
        ["Quy mô", "Cách mở rộng", "Rủi ro cần kiểm"],
        [
            ("2-4 zone", "Một Central Controller, một bridge.", "Dòng relay/nguồn valve, nhiễu flow pulse."),
            ("5-10 zone", "Một hoặc hai Central Controller, chia nhóm manifold.", "Channel ESP-NOW, ACK timeout, áp lực bơm."),
            ("10-20 zone", "Nhiều controller 4-6 zone/cụm; chia `maxConcurrentGroup`.", "Không phụ thuộc một board duy nhất; cần mapping group rõ."),
            (">20 zone", "Cân nhắc RS485/MQTT local hoặc nhiều gateway.", "Ngoài scope đồ án nhưng contract CoreIoT vẫn giữ được."),
        ],
        widths=[1500, 4150, 3710],
        font_size=8.8,
    )

    add_heading(doc, "7. Logic điều khiển, Irrigation Scheduler và mode vận hành", 1)
    add_heading(doc, "7.1. Control mode và System operation mode", 2)
    add_table(
        doc,
        ["Lớp mode", "Giá trị", "Ý nghĩa"],
        [
            ("Control mode theo field", "MANUAL / AUTO / DISABLED", "Người vận hành quyết định field được auto hay chỉ manual; DISABLED khóa tưới thường lệ."),
            ("System operation mode", "NORMAL / DEGRADED / SAFE-IDLE", "Trạng thái vận hành toàn hệ dựa trên cloud, safety, reboot và lỗi nghiêm trọng."),
            ("Manual ON", "TTL / maxWater bắt buộc", "Không cho ON vô thời hạn; vẫn phải qua hard interlock."),
            ("Manual OFF", "Latch theo thời gian", "Chặn Auto bật lại ngay sau khi người dùng vừa tắt."),
        ],
        widths=[2500, 2300, 4560],
        font_size=8.8,
    )
    add_heading(doc, "7.2. Irrigation Scheduler", 2)
    add_numbers(
        doc,
        [
            "Đọc danh sách zone cần tưới: moisture dưới min/target, VPD cao, chưa bị DISABLED hoặc safety lock.",
            "Tính điểm ưu tiên theo độ khô, VPD/airTemp, lần tưới gần nhất và crop/soil preset.",
            "Chọn tối đa `maxConcurrentZones` hoặc `maxConcurrentGroup` theo năng lực bơm/đường ống.",
            "Mở valve từng zone đã chọn, bật main pump khi có ít nhất một valve active.",
            "Theo dõi moisture, flow, maxDuration, maxWaterPerCycle, tankLowSwitch và WaterloggingRisk.",
            "Đóng riêng zone đạt target hoặc bị block; tắt main pump khi không còn valve active.",
            "Ghi log quyết định và publish trạng thái để dashboard giải thích được hành vi scheduler.",
        ],
    )
    add_heading(doc, "7.3. Logic tưới đa biến", 2)
    add_table(
        doc,
        ["Tác vụ", "Điều kiện BẬT", "Điều kiện TẮT / khóa"],
        [
            ("Tưới zone", "moisture < min/target và VPD/airTemp/lịch cho phép; zone không bị lock.", "moisture đạt target/max, maxWater/maxDuration, flow lỗi, waterlogging, manual OFF, hard interlock."),
            ("Quạt", "RH > 80% hoặc airTemp > 32 °C.", "RH/T giảm dưới ngưỡng có hysteresis hoặc mode DISABLED."),
            ("Đèn", "lightLux < 200 trong giờ canh tác.", "Ngoài lịch hoặc lux đủ; ưu tiên Could/Should."),
            ("Drain hook", "WaterloggingRisk hoặc lệnh manual drain.", "Future: chưa bắt buộc chạy bơm thoát thật trong lõi."),
        ],
        widths=[1600, 3900, 3860],
        font_size=8.8,
    )
    add_heading(doc, "7.4. Manual override", 2)
    add_table(
        doc,
        ["Case", "Quy tắc"],
        [
            ("Manual ON", "Chỉ nhận khi commandId mới, TTL còn hiệu lực, control mode cho phép và không vi phạm hard interlock."),
            ("Manual OFF", "Luôn được ưu tiên hơn manual ON/auto; tạo latch để auto không bật lại trong khoảng cấu hình."),
            ("Emergency OFF", "Ưu tiên ngay sau hard interlock; tắt toàn bộ actuator liên quan."),
            ("Lệnh cũ / trùng", "Reject hoặc idempotent theo commandId; không thực thi lại nếu đã ACK thành công."),
            ("Mất ACK từ Central", "Retry giới hạn; quá ngưỡng chuyển trạng thái UNKNOWN/SAFE-IDLE tùy mức rủi ro."),
        ],
        widths=[2100, 7260],
        font_size=8.8,
    )
    add_callout(
        doc,
        "Quyết định #6 — Scheduler ở edge, cloud vẫn là mặt điều khiển",
        "CoreIoT cung cấp dashboard, attribute, rule chain và RPC; nhưng scheduler cần chạy được ở Pi để DEGRADED mode vẫn tưới tối giản khi mất cloud. "
        "Cloud có thể đề xuất/ra lệnh, Pi là chốt kiểm safety cuối trước khi xuống Central Controller.",
    )

    add_heading(doc, "8. Cảnh báo, WaterloggingRisk và drainage hook", 1)
    add_heading(doc, "8.1. Lớp alarm", 2)
    add_table(
        doc,
        ["Lớp alarm", "Ví dụ", "Clear condition"],
        [
            ("Template lõi", "Low/High Moisture, Low Battery.", "Theo rule template và ngưỡng propagate."),
            ("Môi trường", "High Temperature, Low Water Level hiển thị.", "Hết ngưỡng liên tục N mẫu."),
            ("Kết nối", "Node Offline, Central Offline, seqGap cao.", "Có telemetry hợp lệ trở lại."),
            ("Thủy lực", "PumpDryRun, ZoneFlowLow, ValveLeak/LeakSuspected.", "Flow và state trở lại hợp lệ; cần debounce."),
            ("Safety", "TankLowCritical, WaterQuotaExceeded, WaterloggingRisk.", "Clear thủ công hoặc tự clear khi dưới ngưỡng an toàn đủ lâu."),
            ("Data quality", "SensorAnomaly, invalid/stale telemetry.", "Có dữ liệu hợp lệ trong cửa sổ quan sát."),
        ],
        widths=[1800, 4050, 3510],
        font_size=8.8,
    )
    add_heading(doc, "8.2. WaterloggingRisk - phần “tiêu” tối thiểu", 2)
    add_para(
        doc,
        "Trong phạm vi v2.2, “tiêu” không có nghĩa bắt buộc phải bơm thoát nước chủ động. Phần tối thiểu là phát hiện zone quá ẩm/nguy cơ úng, khóa tưới zone đó và cảnh báo rõ ràng. "
        "Điều kiện cơ bản: `moisture > floodMoistureThreshold` trong N mẫu liên tiếp, hoặc `moisture > maxMoistureThreshold` kéo dài sau khi đã dừng tưới. "
        "Clear condition đề xuất: moisture xuống dưới `maxMoistureThreshold - hysteresis` trong M mẫu và không còn flow bất thường.",
    )
    add_table(
        doc,
        ["Tình huống", "Hành động"],
        [
            ("Zone chưa tưới nhưng quá ẩm", "Tạo `WaterloggingRisk`, khóa tưới zone, hiển thị cảnh báo."),
            ("Zone đang tưới và vượt flood threshold", "Đóng valve zone ngay, cập nhật reason `WATERLOGGING_LOCK`."),
            ("Nhiều sensor trong một zone lệch nhau", "Ưu tiên max/percentile hoặc sensor-level alarm, không chỉ dùng trung bình."),
            ("Manual ON khi đang WaterloggingRisk", "Reject trừ khi admin override được thiết kế riêng; mặc định không vượt safety lock."),
        ],
        widths=[3200, 6160],
        font_size=8.8,
    )
    add_heading(doc, "8.3. Drain Pump/Valve Controller - thiết kế mở", 2)
    add_table(
        doc,
        ["Telemetry/RPC đề xuất", "Ý nghĩa"],
        [
            ("drainPumpState / drainValveState", "Trạng thái actuator tiêu nước nếu có phần cứng."),
            ("drainMode", "MANUAL/AUTO/DISABLED tương tự control mode."),
            ("Drain_ON / Drain_OFF", "RPC mở rộng, chưa bắt buộc trong lõi."),
            ("drainReason", "WATERLOGGING, MANUAL, TEST."),
            ("drainMaxDuration", "Giới hạn an toàn nếu triển khai thật."),
        ],
        widths=[3000, 6360],
        font_size=8.8,
    )

    add_heading(doc, "9. Thiết kế phần mềm Raspberry Pi gateway", 1)
    add_para(
        doc,
        "Raspberry Pi 4 chạy Raspberry Pi OS 64-bit và đóng vai edge gateway đa module. Thiết kế phần mềm ưu tiên interface rõ, cấu hình ngoài mã, log đầy đủ và khả năng thay SIM bằng phần cứng thật mà không đổi contract CoreIoT.",
    )
    add_table(
        doc,
        ["Module", "Chức năng", "Vào / Ra"],
        [
            ("Sensor Data Receiver", "Đọc UART từ ESP32 bridge, kiểm CRC, đưa message vào queue.", "UART frame -> message sạch."),
            ("Sensor Reader / Parser", "Tách MAC, node_type, schema, sequence, payload.", "Raw message -> typed event."),
            ("Flow Counter Aggregator", "Nhận pulseCounter từng nhánh hoặc SIM mode; tính delta, flow rate, daily volume.", "Central telemetry -> zone water stats."),
            ("Data Mapper + Quality Gate", "Map key lõi/mở rộng; kiểm range, type, rate-of-change, dedup; publish quality/status.", "Typed event -> telemetry hợp lệ."),
            ("CoreIoT Telemetry Client", "MQTT 3.1.1 QoS 1, keep-alive, LWT, token per logical device.", "Telemetry -> CoreIoT."),
            ("RPC Listener", "Subscribe RPC, whitelist method, kiểm commandId/timestamp/source/TTL.", "Cloud command -> validated command."),
            ("Command Dispatcher / Actuator Coordinator", "Chuyển lệnh xuống Central qua ESP-NOW; direct GPIO chỉ dùng lab/demo gần Pi.", "Validated command -> downlink + ACK."),
            ("Irrigation Scheduler", "Chọn zone, giới hạn maxConcurrentZones, ưu tiên khô/VPD/lần tưới.", "State/cache -> actuator plan."),
            ("Local Safety Handler & Mode Manager", "Watchdog, dry-run, tankLow, quota, waterlogging lock, NORMAL/DEGRADED/SAFE-IDLE.", "State -> allow/block/reason."),
            ("Anomaly Detector", "Đọc telemetry sạch + quality events, phát SensorAnomaly/LeakSuspected/ValveLeak.", "Telemetry window -> alarm candidate."),
            ("Logger / Replay Buffer", "Log xoay vòng, lưu timestamp gốc tại Pi, replay khi MQTT phục hồi.", "Events -> file buffer + replay."),
        ],
        widths=[2200, 4450, 2710],
        font_size=8.2,
    )
    add_heading(doc, "9.1. Cấu hình chính", 2)
    add_table(
        doc,
        ["File / nhóm config", "Nội dung"],
        [
            ("nodeMap.yaml", "MAC -> node id, zone id, logical device token, node_type, enabled."),
            ("fieldConfig.yaml", "min/target/max/flood moisture, crop/soil preset, maxWaterPerCycle/day."),
            ("safety.yaml", "watchdog seconds, dry-run timeout, low-flow threshold, mode transition timeout."),
            ("mqtt.env", "MQTT_HOST, port, token path, QoS, keep-alive, TLS tùy chọn."),
            ("scheduler.yaml", "maxConcurrentZones, priority weights, manual TTL, OFF latch duration."),
            ("quality.yaml", "Range/rate rule theo key, flatline window theo từng key, publish quality/status."),
        ],
        widths=[2500, 6860],
        font_size=8.8,
    )
    add_callout(
        doc,
        "Quyết định #7 — SIM/HIL là first-class data source",
        "Vòng đầu có thể mô phỏng branch flow, sensor và command ACK bằng SIM/HIL, nhưng SIM phải sinh đúng schema như phần cứng thật. "
        "Nhờ vậy thay YF-S201/Central Controller thật chỉ đổi driver/source, không đổi CoreIoT mapping và test case.",
    )

    add_heading(doc, "10. Module phát hiện bất thường", 1)
    add_para(
        doc,
        "Anomaly Detector là module độc lập, chạy ngoài vòng điều khiển an toàn. Nó không tự bật/tắt actuator; nó chỉ phát alarm hoặc gợi ý kiểm tra. "
        "Các hard interlock như bể cạn, watchdog, quota và waterlogging lock phải hoạt động được dù module ML tắt.",
    )
    add_table(
        doc,
        ["Loại bất thường", "Dấu hiệu", "Phương pháp / alarm"],
        [
            ("SensorAnomaly", "Out-of-range, spike, drift, stuck/flatline theo config.", "Range rule + z-score/EWMA cửa sổ trượt; alarm SensorAnomaly."),
            ("LeakSuspected / ValveLeak", "Valve Zi OFF nhưng pulseCounter_Zi vẫn tăng.", "Tương quan valve state với branch pulse; alarm ValveLeak/LeakSuspected."),
            ("ZoneFlowLow", "Valve Zi ON nhưng flow_Zi thấp hoặc không tăng pulse.", "Rule thủy lực; có thể safety block zone."),
            ("PumpDryRun", "Main pump ON nhưng tổng branch flow = 0 trong timeout.", "Safety rule, không chờ ML."),
        ],
        widths=[2200, 3500, 3660],
        font_size=8.8,
    )
    add_heading(doc, "10.1. Flatline theo config", 2)
    add_para(
        doc,
        "Không áp dụng flatline chung cho mọi key. `moisture`, `airTemp`, `airHumidity` có thể kiểm flatline; `battery` thay đổi chậm, `lightLux` ban đêm có thể gần hằng số, phao cạn là tín hiệu nhị phân nên cần cửa sổ riêng hoặc tắt flatline.",
    )
    add_heading(doc, "10.2. Kịch bản demo", 2)
    add_table(
        doc,
        ["Kịch bản", "Cách tiêm lỗi", "Kết quả mong đợi"],
        [
            ("Cảm biến rút dây", "Node gửi status lỗi/không có dữ liệu hợp lệ.", "Không publish key chính; alarm quality/SensorAnomaly."),
            ("Spike moisture", "Bơm giá trị nhảy phi vật lý.", "Drop/clamp theo config; lastInvalid hiển thị."),
            ("Valve rò", "Valve OFF nhưng pulseCounter_Zi tăng.", "LeakSuspected/ValveLeak."),
            ("Zone flow thấp", "Valve ON nhưng flow_Zi thấp.", "ZoneFlowLow; đóng riêng zone nếu vượt timeout."),
        ],
        widths=[2300, 3300, 3760],
        font_size=8.8,
    )

    add_heading(doc, "11. Thiết kế an toàn", 1)
    add_heading(doc, "11.1. Thứ tự ưu tiên", 2)
    add_callout(
        doc,
        "Safety priority cố định",
        "hard interlock > emergency/manual OFF > manual ON có TTL > auto cloud > local rule DEGRADED > mặc định OFF. "
        "Không lệnh nào được vượt bể cạn, quota cứng, watchdog, waterlogging lock hoặc relay fail-safe.",
        fill=WARN_FILL,
    )
    add_heading(doc, "11.2. Ba chế độ vận hành", 2)
    add_table(
        doc,
        ["Mode", "Kích hoạt", "Pi / gateway làm gì", "Actuator"],
        [
            ("NORMAL", "MQTT/cloud ổn định, không lỗi safety nghiêm trọng.", "Cloud rule + scheduler/edge safety phối hợp; Pi thực thi RPC và publish telemetry.", "Theo cloud/scheduler sau safety gate."),
            ("DEGRADED", "Mất MQTT > 60 s hoặc cloud không phản hồi nhưng edge còn sensor/cache.", "Local Control Engine dùng cache min/target/max/quota để tưới tối giản.", "Có thể ON/OFF cục bộ trong giới hạn cứng."),
            ("SAFE-IDLE", "Reboot chưa sync, lỗi nghiêm trọng, mất sensor quan trọng, Central mất ACK kéo dài.", "Dừng tự động, yêu cầu kiểm tra/khôi phục.", "Tất cả actuator OFF."),
        ],
        widths=[1500, 2900, 3650, 1310],
        font_size=8.4,
    )
    add_heading(doc, "11.3. Hard interlock", 2)
    add_table(
        doc,
        ["Interlock", "Điều kiện", "Hành động", "Nơi thực thi"],
        [
            ("Bể cạn", "tankLowSwitch=true hoặc phao báo cạn.", "Cấm bật main pump tuyệt đối; ép OFF nếu đang chạy; alarm Critical.", "Central Controller + Pi."),
            ("Watchdog tưới", "Main pump/valve ON quá maxDuration.", "Đóng zone/pump; log reason; alarm.", "Pi + Central timeout phụ."),
            ("Dry-run", "Main pump ON nhưng tổng branch flow = 0.", "OFF pump; alarm PumpDryRun.", "Pi/Central."),
            ("ZoneFlowLow", "Valve Zi ON nhưng flow_Zi thấp.", "Đóng Zi; giữ zone khác nếu an toàn.", "Pi/Central."),
            ("Hạn mức nước", "Volume/zone/cycle/day vượt trần.", "Khóa tưới zone; alarm.", "Pi."),
            ("Waterlogging lock", "moisture vượt flood/max kéo dài.", "Khóa tưới zone, đóng valve nếu đang tưới.", "Pi + rule chain."),
            ("Relay thường-hở", "Mất điện Pi/Central/relay coil.", "Relay nhả, tải OFF.", "Phần cứng."),
        ],
        widths=[1800, 3000, 3100, 1460],
        font_size=8.3,
    )
    add_heading(doc, "11.4. Cổng kiểm tra dữ liệu nhiều tầng", 2)
    add_table(
        doc,
        ["Tầng", "Kiểm tra", "Xử lý"],
        [
            ("Node ESP32", "Dải thô, lỗi I2C/CRC sensor, trạng thái pin.", "Không gửi rác hoặc gắn status lỗi."),
            ("Bridge/Pi parser", "CRC frame, MAC whitelist, schema, seq duplicate/gap.", "Drop frame lỗi, log seqGap."),
            ("Gateway Quality Gate", "Range, type, rate-of-change, dedup, monotonic counter.", "Không publish key chính; publish quality/status/lastInvalid."),
            ("CoreIoT rule/alarm", "Alarm có debounce, clear condition, dedup entity.", "Không bão alarm; giữ chart chính sạch."),
        ],
        widths=[1900, 4100, 3360],
        font_size=8.8,
    )
    add_table(
        doc,
        ["Key", "Dải hợp lệ", "Quy tắc bổ sung"],
        [
            ("moisture", "0-100", "Nhảy quá ~30 đơn vị/mẫu cần flag; flood threshold dùng cho WaterloggingRisk."),
            ("airTemp", "-10-60 °C", "Flatline theo cửa sổ config; spike bị flag."),
            ("airHumidity", "0-100 %", "Kiểm rate-of-change hợp lý."),
            ("lightLux", "0-100000", "Không flatline ban đêm nếu config tắt."),
            ("pulseCounter", "Monotonic", "Không giảm; delta quá lớn cần flag."),
            ("battery", "0-100", "Không kiểm flatline ngắn."),
            ("tankLowSwitch", "boolean", "Không dùng đo liên tục để thay thế interlock."),
        ],
        widths=[1900, 2100, 5360],
        font_size=8.8,
    )
    add_heading(doc, "11.5. Ma trận phòng bị chéo", 2)
    add_table(
        doc,
        ["Rủi ro", "Cơ chế chính", "Dự phòng / tuyến cuối"],
        [
            ("Tưới tràn", "Dừng theo target/max moisture.", "Quota nước, watchdog, WaterloggingRisk, relay thường-hở."),
            ("Cháy bơm vì cạn nước", "tankLowSwitch hard interlock.", "Dry-run flow = 0, alarm Critical."),
            ("Valve rò", "ValveLeak từ pulseCounter_Zi khi valve OFF.", "Manual inspection, khóa zone."),
            ("Cloud mất", "DEGRADED local control.", "SAFE-IDLE nếu thiếu sensor/cache quan trọng."),
            ("Dữ liệu rác", "Quality gate node + gateway.", "Không publish key chính; alarm có lastInvalid."),
            ("Lệnh cũ/trùng", "commandId + TTL + ACK idempotent.", "Reject stale command, log source."),
            ("Mất điện", "Relay thường-hở.", "Actuator OFF vật lý."),
        ],
        widths=[2200, 3600, 3560],
        font_size=8.8,
    )

    add_heading(doc, "12. Dashboard giám sát và vận hành", 1)
    add_table(
        doc,
        ["Khu vực dashboard", "Nội dung nên có"],
        [
            ("Main state", "Danh sách field/zone, moisture, mode, active irrigation, alarm summary."),
            ("Field state", "Biểu đồ moisture/T/RH/lux/VPD, waterConsumption theo ngày/chu kỳ, trạng thái valve/pump."),
            ("Scheduler panel", "maxConcurrentZones, queue zone, priority score, reason ON/OFF."),
            ("Safety panel", "tankLowSwitch, quota, watchdog, dry-run, WaterloggingRisk, last safety block reason."),
            ("Central Controller panel", "online/offline, ACK status, pumpState, valveState_Zi, flow_Zi, fault flags."),
            ("Quality panel", "quality/status/lastInvalid, seqGap, stale telemetry."),
            ("Manual controls", "MANUAL/AUTO/DISABLED, Manual ON TTL, Manual OFF latch, Emergency OFF."),
        ],
        widths=[2600, 6760],
        font_size=8.8,
    )
    add_callout(
        doc,
        "Nguyên tắc dashboard",
        "Dashboard không chỉ có nút bật/tắt. Nó phải giải thích vì sao hệ tưới hoặc không tưới: mode, safety block, scheduler reason, quality status và alarm clear condition.",
    )

    add_heading(doc, "13. Thiết kế phần cứng", 1)
    add_heading(doc, "13.1. Danh mục theo vai trò", 2)
    add_table(
        doc,
        ["Khối", "Linh kiện đề xuất", "Ghi chú"],
        [
            ("Sensor node", "ESP32-WROOM-32/ESP32-S3, capacitive soil moisture, SHT31/DHT22, BH1750.", "Mỗi zone một cụm; pin/RSSI/fw_version."),
            ("Bridge/receiver", "ESP32 ESP-NOW bridge, UART sang Pi.", "Có uplink và downlink; giữ channel cố định."),
            ("Gateway", "Raspberry Pi 4 Model B, Raspberry Pi OS 64-bit.", "Chạy 11 module, MQTT, log/replay."),
            ("Central/Manifold", "ESP32, relay/driver, input flow pulse, low-water switch.", "Quản lý 4-6 zone/cụm."),
            ("Flow branch", "YF-S201 hoặc flow sensor Hall tương đương.", "Hiệu chuẩn pulsesPerLiter theo sensor."),
            ("Valve", "Solenoid valve hoặc valve DC phù hợp áp lực.", "Cần driver/relay và diode flyback nếu tải cảm."),
            ("Main pump", "Bơm DC/AC theo mô hình lab.", "Không drive trực tiếp từ GPIO."),
            ("Relay/driver", "Opto-isolated relay/MOSFET driver, tách JD-VCC.", "Relay thường-hở để mất điện = OFF."),
            ("Tank sensor", "Phao cạn Must; HC-SR04/analog level Should.", "Phao cạn là interlock, đo liên tục chỉ hiển thị/ước lượng."),
            ("ADC ngoài", "ADS1115 nếu analog nhiễu.", "Dùng cho soil/pH future/level analog."),
            ("pH future", "Probe pH analog + board chuyên dụng.", "Không nằm lõi; cần hiệu chuẩn 2-3 điểm và bảo trì."),
        ],
        widths=[1900, 3600, 3860],
        font_size=8.2,
    )
    add_heading(doc, "13.2. Nguyên tắc điện", 2)
    add_bullets(
        doc,
        [
            "GPIO Pi/ESP32 không 5 V tolerant; tín hiệu 5 V phải qua level shifter hoặc opto.",
            "Không cấp tải relay/motor/valve trực tiếp từ GPIO; tách nguồn logic và nguồn tải.",
            "Tải cảm như bơm, valve, quạt phải có diode flyback/snubber phù hợp.",
            "Relay/driver ưu tiên thường-hở hoặc fail de-energized: mất điện nghĩa là OFF.",
            "Dây flow pulse, low-water switch và relay nên ngắn, đặt gần Central Controller.",
            "Nguồn bơm/valve cần đủ dòng khởi động; Pi và ESP32 dùng nguồn riêng ổn định.",
        ],
    )
    add_heading(doc, "13.3. SIM/HIL và thay thế linh kiện", 2)
    add_table(
        doc,
        ["Mức kiểm chứng", "Dùng để kiểm gì"],
        [
            ("SIM-1 publish CoreIoT", "Data contract, dashboard, calculated field, alarm cơ bản."),
            ("SIM-2 scheduler", "maxConcurrentZones, quota, manual TTL/latch, WaterloggingRisk."),
            ("HIL-1 Central fake", "Pi gửi command, nhận ACK, trạng thái valve/pump giả lập."),
            ("HIL-2 Flow pulse thật", "Đọc pulse YF-S201, dry-run/ZoneFlowLow/ValveLeak."),
            ("Full hardware", "Phao cạn, relay thường-hở, bơm/valve thật, fail-safe mất điện."),
        ],
        widths=[2500, 6860],
        font_size=8.8,
    )
    add_heading(doc, "13.4. OTA - thiết kế mở", 2)
    add_para(
        doc,
        "Trong phạm vi đồ án, cập nhật firmware node/controller qua USB là đủ. Tuy nhiên firmware ESP32 nên bật OTA partition, gắn `fw_version`, ghi rollback path và tách config để sau này thêm OTA hybrid mà không đổi data contract.",
    )

    add_heading(doc, "14. Kiểm thử, nghiệm thu và lộ trình", 1)
    add_heading(doc, "14.1. Nhóm kiểm thử chính", 2)
    add_table(
        doc,
        ["Nhóm", "Case quan trọng"],
        [
            ("Data contract", "TC-01..TC-09: moisture, pulseCounter, battery, dashboard/lịch lõi, calculated field và alarm template."),
            ("Môi trường mở rộng", "TC-10..TC-12: T/RH, lux, tankLowSwitch/waterLevel."),
            ("Truyền thông", "TC-13..TC-15: ESP-NOW, UART CRC, MAC mapping."),
            ("Flow/actuator", "TC-16..TC-19: branch flow/YF-S201, pump/valve/fan/light RPC."),
            ("Alarm/rule", "TC-20..TC-24: high temp, node offline, VPD/RH, lux/lịch."),
            ("Safety", "TC-25..TC-28, TC-34..TC-37: watchdog, dry-run, cloud loss, reboot, quota, quality, relay fail-safe, anti-storm."),
            ("Anomaly", "TC-31..TC-32: SensorAnomaly, ValveLeak/LeakSuspected."),
            ("v2.2 topology", "TC-38..TC-45: main pump + valve/flow theo zone, scheduler, Central ACK, WaterloggingRisk, drain hook, crop preset, mode split, manual TTL/latch."),
        ],
        widths=[2100, 7260],
        font_size=8.8,
    )
    add_callout(
        doc,
        "Lưu ý truy vết TC-09",
        "Để tránh lỗ truy vết từ SRS draft, bản thiết kế này gán TC-09 cho kiểm dashboard/lịch lõi của Smart Irrigation Template: field state, inclusion schedule và thao tác vận hành cơ bản.",
        fill=WARN_FILL,
    )
    add_heading(doc, "14.2. Lộ trình triển khai đề xuất", 2)
    add_table(
        doc,
        ["Giai đoạn", "Mục tiêu", "Kết quả ra"],
        [
            ("P0 - CoreIoT skeleton", "Tạo entity/profile/rule/dashboard lõi và mapping token.", "Dashboard nhận telemetry SIM, alarm lõi chạy."),
            ("P1 - Sensor uplink", "ESP-NOW sensor node -> bridge -> Pi -> CoreIoT.", "Moisture/T/RH/lux theo zone, quality gate cơ bản."),
            ("P2 - Water SIM", "Mô phỏng Central, branch pulse, pump/valve state.", "Scheduler và water statistic theo zone chạy được."),
            ("P3 - Central HIL", "ESP32 Central thật nhận command/ACK, điều khiển relay giả/tải nhẹ.", "RPC xuống hiện trường ổn định."),
            ("P4 - Safety full", "Phao cạn, dry-run, quota, watchdog, waterlogging, DEGRADED.", "Test safety pass trước khi demo auto."),
            ("P5 - Hardware integration", "Flow YF-S201/valve/pump thật trong lab.", "Demo end-to-end, log và báo cáo nghiệm thu."),
        ],
        widths=[1900, 4100, 3360],
        font_size=8.8,
    )
    add_heading(doc, "14.3. Mười quyết định thiết kế cốt lõi", 2)
    add_table(
        doc,
        ["#", "Quyết định", "Tóm tắt"],
        [
            ("1", "v2.2 đổi topology nước", "Main pump + manifold + valve/flow theo zone."),
            ("2", "Phân vai bốn tầng", "Sensor node, Central Controller, Pi gateway, CoreIoT."),
            ("3", "ESP-NOW cho node và Central", "Trễ thấp, ít hạ tầng, phù hợp demo field."),
            ("4", "Logical device theo zone", "Physical controller có thể gom zone, cloud vẫn tách thống kê/alarm."),
            ("5", "Giữ key lõi khi publish", "`pulseCounter_Zi` nội bộ map thành `pulseCounter` trên meter của zone."),
            ("6", "Scheduler ở edge", "DEGRADED vẫn tưới tối giản bằng cache."),
            ("7", "SIM/HIL first-class", "Thay phần cứng thật không đổi interface."),
            ("8", "Anomaly ngoài vòng safety", "ML chỉ cảnh báo, không tự điều khiển actuator."),
            ("9", "Safety priority cố định", "Hard interlock luôn thắng manual/cloud/auto."),
            ("10", "pH/drainage/OTA là mở rộng", "Không chặn lõi tưới tiêu v2.2."),
        ],
        widths=[700, 2800, 5860],
        font_size=8.8,
    )
    add_heading(doc, "14.4. Definition of Done", 2)
    add_bullets(
        doc,
        [
            "CoreIoT có đủ entity lõi/mở rộng, mapping theo zone và dashboard giải thích được trạng thái.",
            "Pi publish đúng key lõi và key mở rộng, không đưa dữ liệu rác vào chart chính.",
            "Central/Manifold Controller nhận lệnh, ACK, gửi pump/valve/flow/low-water telemetry.",
            "Scheduler không vượt `maxConcurrentZones`, đóng riêng zone khi đạt target/hạn mức/lỗi flow.",
            "Mất cloud chuyển DEGRADED, nối lại chuyển NORMAL và replay buffer hợp lệ.",
            "Bể cạn, watchdog, dry-run, quota, WaterloggingRisk và relay fail-safe đều pass test.",
            "Anomaly Detector phát hiện SensorAnomaly và ValveLeak/LeakSuspected mà không tự điều khiển actuator.",
            "Tài liệu SRS, thiết kế hệ thống, README triển khai và kịch bản test thống nhất mã FR/NFR/TC.",
        ],
    )
    add_callout(
        doc,
        "Kết luận thiết kế",
        "Bản v2.2 giữ giá trị cốt lõi của v2.1 nhưng đưa hệ thống gần hiện trường hơn: đo theo zone, điều khiển nước tại trạm, scheduler có giới hạn thủy lực, safety nhiều lớp và khả năng tự chủ khi mất cloud. "
        "Đây là baseline phù hợp để chuyển sang thiết kế chi tiết, code gateway, firmware ESP32 và bộ test nghiệm thu.",
    )

    doc.core_properties.title = "Thiết kế hệ thống v2.2 SmartFarm an toàn"
    doc.core_properties.subject = "Smart Farm CoreIoT system design"
    doc.core_properties.keywords = "CoreIoT, ESP-NOW, Raspberry Pi, Smart Irrigation, Safety, Waterlogging"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT))


if __name__ == "__main__":
    build_doc()
    print(OUTPUT)
